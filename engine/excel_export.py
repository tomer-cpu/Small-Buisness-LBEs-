"""Export financial reports to Excel."""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from models.currency import get_symbol
from engine.financial_engine import ForecastReport, FinancialEngine


HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TOTAL_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
TOTAL_FONT = Font(bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14)
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


def _style_header_row(ws, row, max_col):
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN_BORDER


def _style_data_rows(ws, start_row, end_row, max_col):
    for row in range(start_row, end_row + 1):
        for col in range(1, max_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center")


def _auto_width(ws, max_col):
    for col in range(1, max_col + 1):
        max_len = 0
        letter = get_column_letter(col)
        for cell in ws[letter]:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = max(max_len + 3, 12)


def export_to_excel(report: ForecastReport, engine: FinancialEngine,
                    filepath: str) -> None:
    """Export the full financial report to an Excel file."""
    wb = Workbook()
    sym = get_symbol(report.currency_code)

    # --- Sheet 1: Monthly Summary ---
    ws1 = wb.active
    ws1.title = "Monthly Summary"

    ws1.cell(row=1, column=1, value="Rolling Financial Forecast").font = TITLE_FONT
    ws1.cell(row=2, column=1, value=f"Currency: {report.currency_code} ({sym})")
    ws1.cell(row=3, column=1, value=f"Period: {report.num_months} months")

    headers = ["Month", "Units", f"Revenue ({sym})", f"Cost ({sym})",
               f"Gross Profit ({sym})", "Margin %"]
    for col, h in enumerate(headers, 1):
        ws1.cell(row=5, column=col, value=h)
    _style_header_row(ws1, 5, len(headers))

    row = 6
    for ms in report.month_summaries:
        ws1.cell(row=row, column=1, value=ms.month_label)
        ws1.cell(row=row, column=2, value=ms.total_units)
        ws1.cell(row=row, column=3, value=round(ms.total_revenue, 2))
        ws1.cell(row=row, column=4, value=round(ms.total_cost, 2))
        ws1.cell(row=row, column=5, value=round(ms.gross_profit, 2))
        ws1.cell(row=row, column=6, value=f"{ms.margin_percent:.1f}%")
        row += 1

    # Total row
    total_row = row
    ws1.cell(row=total_row, column=1, value="TOTAL")
    ws1.cell(row=total_row, column=2,
             value=sum(ms.total_units for ms in report.month_summaries))
    ws1.cell(row=total_row, column=3, value=round(report.grand_total_revenue, 2))
    ws1.cell(row=total_row, column=4, value=round(report.grand_total_cost, 2))
    ws1.cell(row=total_row, column=5, value=round(report.grand_total_profit, 2))
    ws1.cell(row=total_row, column=6, value=f"{report.grand_margin_percent:.1f}%")

    for col in range(1, len(headers) + 1):
        cell = ws1.cell(row=total_row, column=col)
        cell.fill = TOTAL_FILL
        cell.font = TOTAL_FONT

    _style_data_rows(ws1, 6, total_row, len(headers))
    _auto_width(ws1, len(headers))

    # --- Sheet 2: Product Summary ---
    ws2 = wb.create_sheet("Product Summary")
    product_summary = engine.get_product_summary(report)

    headers2 = ["SKU", "Product", "Total Units", f"Total Revenue ({sym})",
                f"Total Cost ({sym})", f"Total Profit ({sym})", "Margin %"]
    for col, h in enumerate(headers2, 1):
        ws2.cell(row=1, column=col, value=h)
    _style_header_row(ws2, 1, len(headers2))

    row = 2
    for s in sorted(product_summary.values(),
                    key=lambda x: x["total_profit"], reverse=True):
        ws2.cell(row=row, column=1, value=s["sku"])
        ws2.cell(row=row, column=2, value=s["name"])
        ws2.cell(row=row, column=3, value=s["total_units"])
        ws2.cell(row=row, column=4, value=round(s["total_revenue"], 2))
        ws2.cell(row=row, column=5, value=round(s["total_cost"], 2))
        ws2.cell(row=row, column=6, value=round(s["total_profit"], 2))
        ws2.cell(row=row, column=7, value=f"{s['margin_percent']:.1f}%")
        row += 1

    _style_data_rows(ws2, 2, row - 1, len(headers2))
    _auto_width(ws2, len(headers2))

    # --- Sheet 3: Detailed Breakdown ---
    ws3 = wb.create_sheet("Detail by Month")
    headers3 = ["SKU", "Product", "Month", "Units", f"Unit Price ({sym})",
                f"Revenue ({sym})", f"Cost ({sym})", f"Profit ({sym})", "Margin %"]
    for col, h in enumerate(headers3, 1):
        ws3.cell(row=1, column=col, value=h)
    _style_header_row(ws3, 1, len(headers3))

    row = 2
    for r in report.product_results:
        ws3.cell(row=row, column=1, value=r.sku)
        ws3.cell(row=row, column=2, value=r.product_name)
        ws3.cell(row=row, column=3, value=r.month_label)
        ws3.cell(row=row, column=4, value=r.units)
        ws3.cell(row=row, column=5, value=round(r.selling_price, 2))
        ws3.cell(row=row, column=6, value=round(r.total_revenue, 2))
        ws3.cell(row=row, column=7, value=round(r.total_cost, 2))
        ws3.cell(row=row, column=8, value=round(r.gross_profit, 2))
        ws3.cell(row=row, column=9, value=f"{r.margin_percent:.1f}%")
        row += 1

    _style_data_rows(ws3, 2, row - 1, len(headers3))
    _auto_width(ws3, len(headers3))

    wb.save(filepath)
